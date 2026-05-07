from pathlib import Path
import math
import textwrap

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


OUT = Path("护士节公益健康活动方案_直接汇报版.pptx")
ASSET_DIR = Path("ppt_assets_nurses_day")
FONT = "Microsoft YaHei"

WIDE_W, WIDE_H = 13.333, 7.5

COL = {
    "navy": RGBColor(21, 55, 84),
    "blue": RGBColor(43, 127, 198),
    "blue2": RGBColor(86, 160, 214),
    "teal": RGBColor(66, 184, 164),
    "mint": RGBColor(228, 248, 244),
    "mint2": RGBColor(207, 239, 234),
    "green": RGBColor(111, 190, 123),
    "orange": RGBColor(244, 151, 82),
    "yellow": RGBColor(247, 202, 93),
    "red": RGBColor(231, 102, 96),
    "pink": RGBColor(246, 198, 205),
    "bg": RGBColor(246, 251, 253),
    "white": RGBColor(255, 255, 255),
    "gray": RGBColor(88, 104, 120),
    "light_gray": RGBColor(219, 232, 238),
    "pale_blue": RGBColor(234, 245, 252),
}

PIL = {
    "navy": (21, 55, 84),
    "blue": (43, 127, 198),
    "blue2": (86, 160, 214),
    "teal": (66, 184, 164),
    "mint": (228, 248, 244),
    "mint2": (207, 239, 234),
    "green": (111, 190, 123),
    "orange": (244, 151, 82),
    "yellow": (247, 202, 93),
    "red": (231, 102, 96),
    "pink": (246, 198, 205),
    "bg": (246, 251, 253),
    "white": (255, 255, 255),
    "gray": (88, 104, 120),
    "line": (191, 218, 229),
    "skin": (246, 199, 161),
    "dark": (34, 59, 79),
}


def rgb_tuple(rgb):
    return (rgb[0], rgb[1], rgb[2])


def hi_canvas(w=1200, h=760, scale=2):
    img = Image.new("RGB", (w * scale, h * scale), PIL["bg"])
    return img, ImageDraw.Draw(img), scale


def rr(draw, box, radius, fill, outline=None, width=1):
    draw.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=width)


def ellipse(draw, box, fill, outline=None, width=1):
    draw.ellipse(box, fill=fill, outline=outline, width=width)


def line(draw, points, fill, width=4):
    draw.line(points, fill=fill, width=width, joint="curve")


def draw_person(draw, cx, cy, s=1.0, role="nurse", pose="stand", color=None):
    skin = PIL["skin"]
    hair = PIL["dark"]
    coat = PIL["white"] if role in {"nurse", "doctor"} else (122, 168, 203)
    accent = color or (43, 127, 198)
    # Legs
    leg_y = cy + 110 * s
    if pose == "walk":
        line(draw, [(cx - 20 * s, cy + 90 * s), (cx - 48 * s, leg_y)], PIL["navy"], int(10 * s))
        line(draw, [(cx + 18 * s, cy + 90 * s), (cx + 48 * s, leg_y)], PIL["navy"], int(10 * s))
    else:
        line(draw, [(cx - 18 * s, cy + 85 * s), (cx - 25 * s, leg_y)], PIL["navy"], int(10 * s))
        line(draw, [(cx + 18 * s, cy + 85 * s), (cx + 25 * s, leg_y)], PIL["navy"], int(10 * s))
    # Body
    rr(draw, (cx - 45 * s, cy - 5 * s, cx + 45 * s, cy + 95 * s), int(18 * s), coat, PIL["line"], int(2 * s))
    if role == "nurse":
        rr(draw, (cx - 30 * s, cy - 38 * s, cx + 30 * s, cy - 10 * s), int(8 * s), PIL["white"], PIL["line"], int(2 * s))
        line(draw, [(cx - 8 * s, cy - 25 * s), (cx + 8 * s, cy - 25 * s)], PIL["red"], int(4 * s))
        line(draw, [(cx, cy - 33 * s), (cx, cy - 17 * s)], PIL["red"], int(4 * s))
    elif role == "doctor":
        rr(draw, (cx - 22 * s, cy + 22 * s, cx + 22 * s, cy + 32 * s), int(4 * s), accent)
    else:
        rr(draw, (cx - 35 * s, cy + 8 * s, cx + 35 * s, cy + 70 * s), int(18 * s), accent)
    # Arms
    line(draw, [(cx - 42 * s, cy + 25 * s), (cx - 72 * s, cy + 62 * s)], skin, int(10 * s))
    if pose == "point":
        line(draw, [(cx + 42 * s, cy + 25 * s), (cx + 88 * s, cy + 10 * s)], skin, int(10 * s))
    else:
        line(draw, [(cx + 42 * s, cy + 25 * s), (cx + 70 * s, cy + 63 * s)], skin, int(10 * s))
    # Neck and head
    rr(draw, (cx - 12 * s, cy - 20 * s, cx + 12 * s, cy + 5 * s), int(7 * s), skin)
    ellipse(draw, (cx - 34 * s, cy - 72 * s, cx + 34 * s, cy - 8 * s), skin, None)
    draw.pieslice((cx - 36 * s, cy - 78 * s, cx + 36 * s, cy - 6 * s), 180, 360, fill=hair)
    ellipse(draw, (cx - 12 * s, cy - 46 * s, cx - 6 * s, cy - 40 * s), PIL["dark"])
    ellipse(draw, (cx + 9 * s, cy - 46 * s, cx + 15 * s, cy - 40 * s), PIL["dark"])
    line(draw, [(cx - 10 * s, cy - 25 * s), (cx + 10 * s, cy - 25 * s)], PIL["red"], max(2, int(2 * s)))
    # Shoes
    rr(draw, (cx - 40 * s, leg_y - 5 * s, cx - 12 * s, leg_y + 8 * s), int(5 * s), PIL["navy"])
    rr(draw, (cx + 12 * s, leg_y - 5 * s, cx + 40 * s, leg_y + 8 * s), int(5 * s), PIL["navy"])


def draw_hospital(draw, x, y, w, h):
    rr(draw, (x, y, x + w, y + h), 28, (236, 246, 250), PIL["line"], 3)
    rr(draw, (x + w * 0.38, y + h * 0.16, x + w * 0.62, y + h * 0.82), 16, PIL["white"], PIL["line"], 3)
    rr(draw, (x + w * 0.08, y + h * 0.33, x + w * 0.37, y + h * 0.82), 16, PIL["white"], PIL["line"], 3)
    rr(draw, (x + w * 0.63, y + h * 0.33, x + w * 0.92, y + h * 0.82), 16, PIL["white"], PIL["line"], 3)
    for col in range(3):
        for row in range(3):
            wx = x + w * 0.13 + col * w * 0.09
            wy = y + h * 0.40 + row * h * 0.12
            rr(draw, (wx, wy, wx + w * 0.045, wy + h * 0.045), 4, (193, 224, 237))
            wx2 = x + w * 0.68 + col * w * 0.09
            rr(draw, (wx2, wy, wx2 + w * 0.045, wy + h * 0.045), 4, (193, 224, 237))
    for row in range(4):
        for col in range(2):
            wx = x + w * 0.43 + col * w * 0.09
            wy = y + h * 0.24 + row * h * 0.11
            rr(draw, (wx, wy, wx + w * 0.05, wy + h * 0.045), 4, (193, 224, 237))
    cx = x + w * 0.5
    cy = y + h * 0.08
    line(draw, [(cx - 25, cy), (cx + 25, cy)], PIL["red"], 14)
    line(draw, [(cx, cy - 25), (cx, cy + 25)], PIL["red"], 14)


def save_scaled(img, path, scale=2):
    img = img.resize((img.size[0] // scale, img.size[1] // scale), Image.Resampling.LANCZOS)
    img.save(path)


def asset_cover(path):
    img, d, sc = hi_canvas(1500, 900)
    for i in range(12):
        y = 70 + i * 62
        line(d, [(0, y), (1500, y - 80)], (235, 247, 250), 2 * sc)
    draw_hospital(d, 760 * sc, 70 * sc, 610 * sc, 430 * sc)
    rr(d, (120 * sc, 600 * sc, 1380 * sc, 790 * sc), 36 * sc, PIL["white"], PIL["line"], 3 * sc)
    rr(d, (155 * sc, 630 * sc, 530 * sc, 750 * sc), 28 * sc, (235, 248, 245), None)
    rr(d, (585 * sc, 630 * sc, 940 * sc, 750 * sc), 28 * sc, (235, 245, 252), None)
    rr(d, (995 * sc, 630 * sc, 1340 * sc, 750 * sc), 28 * sc, (255, 246, 237), None)
    draw_person(d, 300 * sc, 525 * sc, 2.0 * sc, "nurse", "point")
    draw_person(d, 600 * sc, 560 * sc, 1.45 * sc, "citizen", "stand", PIL["green"])
    draw_person(d, 745 * sc, 565 * sc, 1.25 * sc, "citizen", "stand", PIL["orange"])
    # Clipboard / service symbols
    rr(d, (220 * sc, 690 * sc, 360 * sc, 760 * sc), 18 * sc, PIL["white"], PIL["line"], 3 * sc)
    line(d, [(250 * sc, 720 * sc), (330 * sc, 720 * sc)], PIL["blue"], 5 * sc)
    line(d, [(250 * sc, 742 * sc), (310 * sc, 742 * sc)], PIL["teal"], 5 * sc)
    # Heart pulse
    pts = [(1080 * sc, 700 * sc), (1120 * sc, 700 * sc), (1148 * sc, 662 * sc), (1180 * sc, 738 * sc), (1212 * sc, 700 * sc), (1275 * sc, 700 * sc)]
    line(d, pts, PIL["red"], 8 * sc)
    save_scaled(img, path, sc)


def asset_checkup(path):
    img, d, sc = hi_canvas(1200, 760)
    rr(d, (80 * sc, 500 * sc, 1120 * sc, 650 * sc), 24 * sc, PIL["white"], PIL["line"], 3 * sc)
    rr(d, (150 * sc, 250 * sc, 1050 * sc, 520 * sc), 28 * sc, (235, 248, 245), PIL["line"], 3 * sc)
    draw_person(d, 350 * sc, 440 * sc, 1.65 * sc, "nurse", "point")
    draw_person(d, 650 * sc, 455 * sc, 1.35 * sc, "citizen", "stand", PIL["orange"])
    # table equipment
    rr(d, (500 * sc, 530 * sc, 860 * sc, 590 * sc), 18 * sc, (221, 238, 247), PIL["line"], 3 * sc)
    rr(d, (535 * sc, 470 * sc, 675 * sc, 535 * sc), 14 * sc, PIL["white"], PIL["line"], 3 * sc)
    ellipse(d, (705 * sc, 470 * sc, 785 * sc, 540 * sc), PIL["mint2"], PIL["line"], 3 * sc)
    line(d, [(730 * sc, 505 * sc), (760 * sc, 485 * sc)], PIL["blue"], 5 * sc)
    line(d, [(555 * sc, 500 * sc), (650 * sc, 500 * sc)], PIL["blue"], 5 * sc)
    # queue
    for x, c in [(925, PIL["green"]), (1010, PIL["blue2"])]:
        draw_person(d, x * sc, 470 * sc, 0.75 * sc, "citizen", "stand", c)
    save_scaled(img, path, sc)


def asset_consultation(path):
    img, d, sc = hi_canvas(1200, 760)
    rr(d, (95 * sc, 120 * sc, 1105 * sc, 640 * sc), 34 * sc, PIL["white"], PIL["line"], 4 * sc)
    rr(d, (175 * sc, 455 * sc, 1015 * sc, 560 * sc), 22 * sc, (226, 240, 249), PIL["line"], 3 * sc)
    draw_person(d, 355 * sc, 390 * sc, 1.4 * sc, "doctor", "point")
    draw_person(d, 760 * sc, 405 * sc, 1.3 * sc, "citizen", "stand", PIL["green"])
    rr(d, (520 * sc, 245 * sc, 675 * sc, 360 * sc), 18 * sc, PIL["mint"], PIL["line"], 3 * sc)
    line(d, [(548 * sc, 285 * sc), (648 * sc, 285 * sc)], PIL["blue"], 5 * sc)
    line(d, [(548 * sc, 315 * sc), (628 * sc, 315 * sc)], PIL["teal"], 5 * sc)
    rr(d, (840 * sc, 240 * sc, 985 * sc, 350 * sc), 18 * sc, (255, 246, 237), PIL["line"], 3 * sc)
    line(d, [(865 * sc, 295 * sc), (900 * sc, 295 * sc), (925 * sc, 260 * sc), (955 * sc, 325 * sc)], PIL["red"], 5 * sc)
    save_scaled(img, path, sc)


def asset_exercise(path):
    img, d, sc = hi_canvas(1200, 760)
    rr(d, (80 * sc, 120 * sc, 1120 * sc, 650 * sc), 34 * sc, (232, 249, 242), PIL["line"], 4 * sc)
    # mats
    for x, y, c in [(220, 545, PIL["blue2"]), (520, 560, PIL["green"]), (815, 545, PIL["orange"])]:
        rr(d, (x * sc, y * sc, (x + 200) * sc, (y + 35) * sc), 18 * sc, c, None)
    draw_person(d, 280 * sc, 425 * sc, 1.25 * sc, "nurse", "point")
    draw_person(d, 590 * sc, 440 * sc, 1.15 * sc, "citizen", "point", PIL["green"])
    draw_person(d, 880 * sc, 430 * sc, 1.1 * sc, "citizen", "point", PIL["orange"])
    # movement arcs
    for cx in [280, 590, 880]:
        d.arc((cx * sc - 95 * sc, 250 * sc, cx * sc + 95 * sc, 420 * sc), 205, 335, fill=PIL["blue"], width=5 * sc)
    save_scaled(img, path, sc)


def asset_tribute(path):
    img, d, sc = hi_canvas(1200, 760)
    rr(d, (100 * sc, 100 * sc, 1100 * sc, 640 * sc), 36 * sc, PIL["white"], PIL["line"], 4 * sc)
    rr(d, (170 * sc, 130 * sc, 1030 * sc, 250 * sc), 24 * sc, (235, 245, 252), None)
    # stage
    rr(d, (200 * sc, 530 * sc, 1000 * sc, 610 * sc), 22 * sc, PIL["blue"], None)
    draw_person(d, 350 * sc, 470 * sc, 1.25 * sc, "nurse", "stand")
    draw_person(d, 510 * sc, 470 * sc, 1.25 * sc, "nurse", "stand")
    draw_person(d, 670 * sc, 470 * sc, 1.25 * sc, "nurse", "stand")
    draw_person(d, 840 * sc, 480 * sc, 1.1 * sc, "citizen", "point", PIL["orange"])
    # flower
    line(d, [(820 * sc, 430 * sc), (760 * sc, 385 * sc)], PIL["green"], 5 * sc)
    for dx, dy in [(755, 380), (775, 372), (760, 360), (742, 365)]:
        ellipse(d, ((dx - 12) * sc, (dy - 12) * sc, (dx + 12) * sc, (dy + 12) * sc), PIL["pink"])
    # blessing cards
    for i in range(9):
        x = 210 + (i % 3) * 75
        y = 160 + (i // 3) * 35
        rr(d, (x * sc, y * sc, (x + 48) * sc, (y + 24) * sc), 6 * sc, (255, 246, 210), PIL["line"], 2 * sc)
    save_scaled(img, path, sc)


def asset_phone(path):
    img, d, sc = hi_canvas(960, 760)
    rr(d, (290 * sc, 80 * sc, 670 * sc, 690 * sc), 52 * sc, PIL["navy"], None)
    rr(d, (318 * sc, 122 * sc, 642 * sc, 642 * sc), 32 * sc, PIL["white"], None)
    rr(d, (350 * sc, 165 * sc, 610 * sc, 245 * sc), 20 * sc, PIL["mint"], None)
    # post images
    rr(d, (350 * sc, 280 * sc, 610 * sc, 430 * sc), 20 * sc, (235, 245, 252), PIL["line"], 2 * sc)
    draw_person(d, 450 * sc, 390 * sc, 0.62 * sc, "nurse", "point")
    draw_person(d, 530 * sc, 400 * sc, 0.52 * sc, "citizen", "stand", PIL["orange"])
    for y, w in [(462, 210), (492, 245), (522, 170)]:
        rr(d, (350 * sc, y * sc, (350 + w) * sc, (y + 16) * sc), 8 * sc, (218, 235, 244), None)
    # hearts and share lines
    ellipse(d, (190 * sc, 200 * sc, 245 * sc, 255 * sc), PIL["pink"])
    ellipse(d, (725 * sc, 360 * sc, 780 * sc, 415 * sc), PIL["mint2"])
    line(d, [(720 * sc, 540 * sc), (820 * sc, 500 * sc), (840 * sc, 580 * sc)], PIL["teal"], 6 * sc)
    ellipse(d, (705 * sc, 525 * sc, 735 * sc, 555 * sc), PIL["teal"])
    ellipse(d, (805 * sc, 485 * sc, 835 * sc, 515 * sc), PIL["blue"])
    ellipse(d, (825 * sc, 565 * sc, 855 * sc, 595 * sc), PIL["orange"])
    save_scaled(img, path, sc)


def asset_stage(path):
    img, d, sc = hi_canvas(1200, 760)
    rr(d, (115 * sc, 110 * sc, 1085 * sc, 635 * sc), 38 * sc, (237, 247, 251), PIL["line"], 4 * sc)
    rr(d, (250 * sc, 160 * sc, 950 * sc, 405 * sc), 30 * sc, PIL["white"], PIL["line"], 4 * sc)
    # Background arcs
    d.arc((300 * sc, 185 * sc, 620 * sc, 405 * sc), 190, 350, fill=PIL["teal"], width=10 * sc)
    d.arc((580 * sc, 185 * sc, 900 * sc, 405 * sc), 190, 350, fill=PIL["blue"], width=10 * sc)
    # Screen and stage
    rr(d, (500 * sc, 205 * sc, 700 * sc, 330 * sc), 14 * sc, (222, 240, 249), PIL["line"], 3 * sc)
    rr(d, (175 * sc, 470 * sc, 1025 * sc, 555 * sc), 24 * sc, PIL["blue"], None)
    rr(d, (165 * sc, 555 * sc, 1035 * sc, 600 * sc), 12 * sc, PIL["navy"], None)
    # Equipment
    for x in [185, 955]:
        rr(d, (x * sc, 410 * sc, (x + 55) * sc, 515 * sc), 12 * sc, PIL["navy"], None)
        ellipse(d, ((x + 12) * sc, 428 * sc, (x + 43) * sc, 459 * sc), PIL["gray"])
    draw_person(d, 365 * sc, 425 * sc, 0.86 * sc, "nurse", "stand")
    draw_person(d, 835 * sc, 430 * sc, 0.82 * sc, "citizen", "stand", PIL["green"])
    save_scaled(img, path, sc)


def asset_venue_map(path):
    img, d, sc = hi_canvas(1200, 760)
    rr(d, (80 * sc, 80 * sc, 1120 * sc, 680 * sc), 34 * sc, PIL["white"], PIL["line"], 4 * sc)
    zones = [
        (120, 120, 330, 255, PIL["blue2"]),
        (360, 120, 650, 255, PIL["mint2"]),
        (680, 120, 1040, 255, (255, 236, 215)),
        (120, 300, 430, 470, (224, 242, 229)),
        (465, 300, 765, 470, (235, 245, 252)),
        (800, 300, 1040, 470, (247, 230, 233)),
        (120, 515, 1040, 635, (238, 246, 250)),
    ]
    for z in zones:
        rr(d, tuple(v * sc if isinstance(v, int) else v for v in z[:4]), 20 * sc, z[4], PIL["line"], 3 * sc)
    # flow arrow
    pts = [(150 * sc, 595 * sc), (300 * sc, 595 * sc), (300 * sc, 280 * sc), (520 * sc, 280 * sc), (520 * sc, 595 * sc), (850 * sc, 595 * sc), (850 * sc, 500 * sc)]
    line(d, pts, PIL["blue"], 9 * sc)
    d.polygon([(850 * sc, 500 * sc), (830 * sc, 535 * sc), (872 * sc, 535 * sc)], fill=PIL["blue"])
    # icon dots
    for x, y, c in [(225, 190, PIL["white"]), (505, 190, PIL["white"]), (860, 190, PIL["white"]), (275, 385, PIL["white"]), (615, 385, PIL["white"]), (920, 385, PIL["white"])]:
        ellipse(d, ((x - 28) * sc, (y - 28) * sc, (x + 28) * sc, (y + 28) * sc), c, PIL["line"], 2 * sc)
    save_scaled(img, path, sc)


def asset_education(path):
    img, d, sc = hi_canvas(1200, 760)
    rr(d, (90 * sc, 110 * sc, 1110 * sc, 640 * sc), 36 * sc, PIL["white"], PIL["line"], 4 * sc)
    colors = [PIL["blue2"], PIL["teal"], PIL["green"], PIL["orange"], PIL["red"]]
    for i, c in enumerate(colors):
        x = 145 + i * 200
        rr(d, (x * sc, 175 * sc, (x + 150) * sc, 565 * sc), 22 * sc, (246, 251, 253), PIL["line"], 3 * sc)
        ellipse(d, ((x + 45) * sc, 220 * sc, (x + 105) * sc, 280 * sc), c)
        for y, w in [(325, 105), (360, 85), (395, 110), (430, 70)]:
            rr(d, ((x + 22) * sc, y * sc, (x + 22 + w) * sc, (y + 13) * sc), 6 * sc, (216, 231, 238), None)
    save_scaled(img, path, sc)


def asset_team(path):
    img, d, sc = hi_canvas(1200, 760)
    rr(d, (90 * sc, 115 * sc, 1110 * sc, 645 * sc), 36 * sc, PIL["white"], PIL["line"], 4 * sc)
    centers = [(600, 230), (310, 390), (490, 390), (710, 390), (890, 390), (410, 540), (600, 540), (790, 540)]
    colors = [PIL["blue"], PIL["teal"], PIL["green"], PIL["orange"], PIL["blue2"], PIL["mint2"], PIL["yellow"], PIL["pink"]]
    for idx, (cx, cy) in enumerate(centers):
        if idx > 0:
            line(d, [(600 * sc, 270 * sc), (cx * sc, (cy - 50) * sc)], PIL["line"], 4 * sc)
        ellipse(d, ((cx - 52) * sc, (cy - 52) * sc, (cx + 52) * sc, (cy + 52) * sc), colors[idx], PIL["white"], 5 * sc)
        draw_person(d, cx * sc, (cy + 25) * sc, 0.42 * sc, "citizen", "stand", PIL["white"])
    save_scaled(img, path, sc)


def asset_metrics(path):
    img, d, sc = hi_canvas(1200, 760)
    rr(d, (100 * sc, 100 * sc, 1100 * sc, 650 * sc), 36 * sc, PIL["white"], PIL["line"], 4 * sc)
    # Dashboard cards
    for i, c in enumerate([PIL["blue"], PIL["teal"], PIL["green"], PIL["orange"]]):
        x = 155 + i * 235
        rr(d, (x * sc, 155 * sc, (x + 180) * sc, 285 * sc), 22 * sc, (246, 251, 253), PIL["line"], 3 * sc)
        ellipse(d, ((x + 25) * sc, 190 * sc, (x + 80) * sc, 245 * sc), c)
        rr(d, ((x + 95) * sc, 190 * sc, (x + 155) * sc, 205 * sc), 7 * sc, (217, 232, 239), None)
        rr(d, ((x + 95) * sc, 225 * sc, (x + 140) * sc, 240 * sc), 7 * sc, (217, 232, 239), None)
    # Chart
    axes = [(190, 560), (190, 360), (965, 560)]
    line(d, [(axes[0][0] * sc, axes[0][1] * sc), (axes[1][0] * sc, axes[1][1] * sc), (axes[2][0] * sc, axes[2][1] * sc)], PIL["line"], 5 * sc)
    bars = [(250, 80, PIL["blue2"]), (360, 130, PIL["teal"]), (470, 105, PIL["green"]), (580, 170, PIL["orange"]), (690, 150, PIL["blue"]), (800, 190, PIL["red"])]
    for x, h, c in bars:
        rr(d, (x * sc, (560 - h) * sc, (x + 62) * sc, 560 * sc), 14 * sc, c, None)
    save_scaled(img, path, sc)


def create_assets():
    ASSET_DIR.mkdir(exist_ok=True)
    assets = {
        "cover": ASSET_DIR / "cover_scene.png",
        "checkup": ASSET_DIR / "checkup_scene.png",
        "consult": ASSET_DIR / "consultation_scene.png",
        "exercise": ASSET_DIR / "exercise_scene.png",
        "tribute": ASSET_DIR / "tribute_scene.png",
        "phone": ASSET_DIR / "social_phone.png",
        "stage": ASSET_DIR / "stage_setup.png",
        "map": ASSET_DIR / "venue_map.png",
        "education": ASSET_DIR / "education_wall.png",
        "team": ASSET_DIR / "team_ops.png",
        "metrics": ASSET_DIR / "metrics_dashboard.png",
    }
    makers = {
        "cover": asset_cover,
        "checkup": asset_checkup,
        "consult": asset_consultation,
        "exercise": asset_exercise,
        "tribute": asset_tribute,
        "phone": asset_phone,
        "stage": asset_stage,
        "map": asset_venue_map,
        "education": asset_education,
        "team": asset_team,
        "metrics": asset_metrics,
    }
    for key, path in assets.items():
        makers[key](path)
    return assets


def fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color


def no_line(shape):
    shape.line.fill.background()


def line_color(shape, color, width=0.8):
    shape.line.color.rgb = color
    shape.line.width = Pt(width)


def textbox(slide, text, x, y, w, h, size=13, color=None, bold=False, align=PP_ALIGN.LEFT,
            valign=MSO_ANCHOR.TOP, margin=0.04, spacing=1.05):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = valign
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, side, Inches(margin))
    for idx, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = line
        p.alignment = align
        p.space_after = Pt(2)
        p.line_spacing = spacing
        for run in p.runs:
            run.font.name = FONT
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = color or COL["navy"]
    return box


def bg(slide):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(WIDE_W), Inches(WIDE_H))
    fill(s, COL["bg"])
    no_line(s)
    band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(WIDE_W), Inches(0.1))
    fill(band, COL["teal"])
    no_line(band)


def title(slide, no, text, subtitle=None, section=None):
    textbox(slide, f"{no:02d}", 0.42, 0.28, 0.55, 0.28, 9, COL["teal"], True, PP_ALIGN.CENTER)
    textbox(slide, text, 0.95, 0.18, 10.9, 0.46, 19, COL["navy"], True)
    if subtitle:
        textbox(slide, subtitle, 0.98, 0.68, 10.7, 0.34, 8.8, COL["gray"])
    if section:
        pill(slide, section, 11.35, 0.28, 1.45, 0.3, COL["mint"], COL["teal"], 7.5)


def footer(slide):
    textbox(slide, "护士节公益健康活动方案｜活动策划 · 视觉设计 · 舞美搭建 · 设备执行一体化", 0.55, 7.12, 6.6, 0.18, 6.6, COL["gray"])
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(7.3), Inches(7.18), Inches(5.45), Inches(0.02))
    fill(line, COL["light_gray"])
    no_line(line)


def pill(slide, text, x, y, w, h, bgc, tc=None, size=9.5, bold=True):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    fill(shape, bgc)
    no_line(shape)
    textbox(slide, text, x + 0.04, y + 0.045, w - 0.08, h - 0.04, size, tc or COL["white"], bold, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, 0.01)


def card(slide, heading, body, x, y, w, h, accent=COL["blue"], body_size=9.8, fill_color=COL["white"]):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    fill(shape, fill_color)
    line_color(shape, COL["light_gray"], 0.8)
    stripe = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(0.08), Inches(h))
    fill(stripe, accent)
    no_line(stripe)
    textbox(slide, heading, x + 0.18, y + 0.12, w - 0.32, 0.28, 10.2, accent, True)
    textbox(slide, body, x + 0.18, y + 0.48, w - 0.32, h - 0.56, body_size, COL["gray"])


def picture(slide, path, x, y, w, h, border=True):
    panel = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    fill(panel, COL["white"])
    line_color(panel, COL["light_gray"], 0.8)
    pic = slide.shapes.add_picture(str(path), Inches(x + 0.06), Inches(y + 0.06), Inches(w - 0.12), Inches(h - 0.12))
    if border:
        return pic
    return pic


def section_label(slide, text, x, y):
    pill(slide, text, x, y, 1.35, 0.32, COL["blue"], COL["white"], 8.2)


def metric_card(slide, number, label, x, y, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(1.58), Inches(0.78))
    fill(shape, COL["white"])
    line_color(shape, COL["light_gray"])
    textbox(slide, number, x + 0.13, y + 0.09, 1.25, 0.25, 15, color, True, PP_ALIGN.CENTER)
    textbox(slide, label, x + 0.08, y + 0.42, 1.42, 0.22, 7.3, COL["gray"], False, PP_ALIGN.CENTER)


def bullet_box(slide, bullets, x, y, w, h, size=10, color=COL["navy"]):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.05)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    for i, line in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.name = FONT
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.space_after = Pt(4)
        p.line_spacing = 1.05
    return box


def icon_circle(slide, label, x, y, color, text_color=COL["white"]):
    c = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(0.46), Inches(0.46))
    fill(c, color)
    no_line(c)
    textbox(slide, label, x, y + 0.08, 0.46, 0.18, 8.2, text_color, True, PP_ALIGN.CENTER)


def slide_1(prs, assets):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    # large visual panel
    s.shapes.add_picture(str(assets["cover"]), Inches(6.35), Inches(0.62), Inches(6.55), Inches(4.05))
    textbox(s, "致敬白衣天使\n共赴健康生活", 0.72, 1.0, 5.55, 1.28, 30, COL["navy"], True, spacing=0.9)
    textbox(s, "医院护士节公益体检与健康生活倡导活动方案", 0.78, 2.54, 5.25, 0.42, 15, COL["blue"], True)
    textbox(
        s,
        "以护士节为契机，面向市民开放公益体检、健康咨询、科学运动互动与健康生活科普，"
        "让护理团队的专业守护被看见，也让市民把健康意识转化为一次真实行动。",
        0.78,
        3.08,
        5.55,
        0.88,
        11.2,
        COL["gray"],
    )
    for i, t in enumerate(["护士节", "公益体检", "健康咨询", "科学运动", "社区公益"]):
        pill(s, t, 0.78 + i * 1.02, 4.28, 0.88, 0.34, [COL["blue"], COL["teal"], COL["green"], COL["orange"], COL["red"]][i], size=7.2)
    card(s, "提案定位", "面向医院甲方的正式提案版本，覆盖活动策略、内容规划、现场设计、宣发引流、舞美物料、人员执行与风险保障。", 0.78, 5.15, 5.35, 0.95, COL["teal"], 9.5)
    card(s, "承办能力", "我们具备策划、设计、美陈搭建、会议物料、展会搭建、自有舞台设备和现场执行团队，可一体化完成落地交付。", 6.62, 5.15, 5.75, 0.95, COL["blue"], 9.5)
    textbox(s, "医院方｜护士节公益健康活动提案", 0.78, 6.72, 5.2, 0.2, 8.5, COL["gray"])


def slide_2(prs, assets):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s); title(s, 2, "汇报路径：先回答为什么做，再回答怎样落地", "本方案按甲方决策逻辑组织，从活动价值、内容体验到执行保障逐层展开。", "目录")
    s.shapes.add_picture(str(assets["team"]), Inches(11.72), Inches(0.72), Inches(0.68), Inches(0.43))
    items = [
        ("01 项目背景", "护士节情感节点、公众健康意识和社区健康管理需求共同构成活动基础。", COL["blue"]),
        ("02 活动策略", "明确活动目标、人群动机、痛点机会和公益服务的核心转化路径。", COL["teal"]),
        ("03 主题创意", "以“被守护的健康”和“守护者的节日”为创意主轴，形成可传播主题。", COL["green"]),
        ("04 内容规划", "设计公益体检、健康咨询、运动互动、护士致敬和科普展陈。", COL["orange"]),
        ("05 现场设计", "规划区域动线、美陈导视、舞美搭建、物料系统与拍照打卡点。", COL["blue2"]),
        ("06 宣发引流", "输出预热、当天、打卡和复盘四类朋友圈文案与传播动作。", COL["red"]),
        ("07 执行保障", "拆分人员职责、设备保障、医护协同、物料管理和风险预案。", COL["teal"]),
        ("08 预期成果", "沉淀公益服务数据、医院品牌声量、社区关系和后续健康管理入口。", COL["green"]),
    ]
    for idx, (h, b, c) in enumerate(items):
        x = 0.74 + (idx % 2) * 6.05
        y = 1.3 + (idx // 2) * 1.35
        card(s, h, b, x, y, 5.55, 1.05, c, 8.8)
    footer(s)


def slide_3(prs, assets):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s); title(s, 3, "项目理解：这不是一场内部节庆，而是一场公益服务开放日", "护士节提供情感入口，公益体检和健康生活倡导提供市民参与理由，医院品牌公益传播完成价值沉淀。", "理解")
    picture(s, assets["checkup"], 0.75, 1.33, 5.15, 3.48)
    textbox(s, "核心判断", 6.25, 1.34, 1.2, 0.25, 10, COL["teal"], True)
    textbox(
        s,
        "本次活动应从“医院内部庆祝”升级为“面向市民的开放式公益健康行动”。"
        "它用护士节的社会情感降低沟通距离，用基础体检和面对面咨询提供真实服务，"
        "再通过朋友圈打卡和社区传播把医院公益形象带到更多人身边。",
        6.25, 1.68, 6.15, 0.98, 12.4, COL["navy"], True,
    )
    layers = [
        ("节日致敬", "让护理团队的专业、耐心和陪伴被公众正式看见。", COL["blue"]),
        ("公益服务", "把基础健康检测和免费咨询带到市民身边，降低参与门槛。", COL["teal"]),
        ("生活倡导", "把指标解释、运动指导和慢病预防转化为可执行建议。", COL["green"]),
        ("品牌传播", "通过真实参与画面建立医院“专业可信、服务有温度”的认知。", COL["orange"]),
    ]
    for i, (h, b, c) in enumerate(layers):
        card(s, h, b, 6.25 + (i % 2) * 3.05, 3.0 + (i // 2) * 1.12, 2.75, 0.92, c, 8.2)
    card(s, "落地原则", "所有现场内容都要形成闭环：市民到场有引导，检测后有人解释，互动后能带走建议，离场前可扫码了解后续服务。", 0.75, 5.15, 11.65, 0.86, COL["teal"], 9.5, COL["mint"])
    footer(s)


def slide_4(prs, assets):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s); title(s, 4, "活动背景：护士节情感与公众健康需求同频", "活动成立的关键，不在于制造热闹，而在于让医院公益服务与市民真实健康需求发生连接。", "背景")
    factors = [
        ("护士节社会情感", "护理团队长期承担照护、沟通、陪伴和应急处置工作。护士节天然具备公众情感认同，适合做一次正式而温暖的社会化表达。", COL["blue"]),
        ("公众健康意识提升", "血压、血糖、体重、睡眠和运动等基础健康话题与生活高度相关。市民愿意参与低门槛、即时获得结果的检测体验。", COL["teal"]),
        ("医院公益服务可感知", "公益价值需要被看见：免费体检、面对面解释、健康建议卡和后续服务二维码，会让市民真正感知医院的服务能力。", COL["green"]),
        ("社区健康管理需求", "中老年、慢病人群、职场亚健康与亲子家庭都需要持续的健康教育和服务入口，医院可借活动建立长期连接。", COL["orange"]),
    ]
    for i, item in enumerate(factors):
        card(s, item[0], item[1], 0.72 + (i % 2) * 6.05, 1.25 + (i // 2) * 1.62, 5.55, 1.25, item[2], 9.2)
    picture(s, assets["consult"], 0.74, 4.75, 4.0, 1.68)
    card(s, "结论", "护士节不仅是致敬护理团队的节点，也是医院把专业健康服务送到市民身边、把公益形象转化为公众体验的最佳窗口。", 4.98, 4.75, 7.42, 1.68, COL["teal"], 10.5, COL["mint"])
    footer(s)


def slide_5(prs, assets):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s); title(s, 5, "痛点与机会：把“知道要健康”变成“今天就行动”", "市民缺少的不是口号，而是低门槛的触发点、看得懂的解释和愿意分享的现场体验。", "策略")
    pairs = [
        ("体检常被拖延", "用公益体检降低行动门槛，现场完成血压、血糖、BMI、心率和基础问诊。"),
        ("指标看不懂", "检测后由医护解释结果，给出复查建议、生活方式建议和健康管理入口。"),
        ("公益活动偏单向", "把讲座、咨询、运动互动、祝福墙和健康打卡组合成可停留的体验。"),
        ("健康传播太官方", "用真实服务照片、亲和朋友圈文案和打卡装置形成市民愿意转发的素材。"),
    ]
    for i, (pain, sol) in enumerate(pairs):
        y = 1.28 + i * 1.18
        pill(s, "痛点", 0.75, y + 0.04, 0.74, 0.29, COL["red"], size=7.4)
        textbox(s, pain, 1.62, y, 2.3, 0.36, 12, COL["navy"], True)
        arrow = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(4.15), Inches(y + 0.05), Inches(0.85), Inches(0.28))
        fill(arrow, COL["light_gray"]); no_line(arrow)
        pill(s, "机会", 5.3, y + 0.04, 0.74, 0.29, COL["teal"], size=7.4)
        textbox(s, sol, 6.15, y - 0.02, 5.95, 0.48, 10.2, COL["gray"])
    picture(s, assets["phone"], 9.8, 4.98, 2.12, 1.55)
    card(s, "策略抓手", "现场体验链路设计为“登记领取流程卡 → 分区体验 → 医护解释 → 健康行动打卡 → 扫码获取后续资料”。这条链路能同时解决参与、理解、传播和留资四件事。", 0.75, 6.0, 8.72, 0.82, COL["blue"], 8.8, COL["white"])
    footer(s)


def slide_6(prs, assets):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s); title(s, 6, "活动目标：五类目标形成公益服务闭环", "目标不是单纯到场人数，而是让公益服务、品牌形象、传播曝光和后续健康管理入口同时建立。", "目标")
    s.shapes.add_picture(str(assets["metrics"]), Inches(11.72), Inches(0.72), Inches(0.68), Inches(0.43))
    goals = [
        ("品牌目标", "让市民通过真实服务感知医院护理团队的专业、耐心与温度，形成对医院公益形象的正向记忆。", COL["blue"]),
        ("公益目标", "为到场市民提供基础检测与初步健康建议，帮助其理解定期体检和基础指标监测的重要性。", COL["teal"]),
        ("传播目标", "设计可拍、可说、可转发的现场内容，让活动不止停留在现场，而能进入朋友圈和社区群。", COL["green"]),
        ("转化目标", "在不商业化打扰的前提下，引导市民了解体检中心、慢病管理、护理门诊和社区服务。", COL["orange"]),
        ("服务目标", "通过清晰动线、人员分工和风险预案，保证老人、亲子和职场人都能顺畅参与。", COL["red"]),
    ]
    for i, (h, b, c) in enumerate(goals):
        x = 0.75 + (i % 3) * 4.05
        y = 1.35 + (i // 3) * 1.72
        card(s, h, b, x, y, 3.62, 1.34, c, 8.3)
    card(s, "建议复盘指标", "到场人数、检测人数、咨询人数、运动互动人数、打卡照片数量、扫码人数、公众号/朋友圈素材数量、满意度反馈和后续预约线索。", 4.8, 4.9, 7.55, 1.02, COL["teal"], 9.6, COL["mint"])
    for i, (num, label, c) in enumerate([("1", "一次健康检测", COL["blue"]), ("1", "一次专业解释", COL["teal"]), ("1", "一个行动建议", COL["green"]), ("1", "一次主动分享", COL["orange"])]):
        metric_card(s, num, label, 0.75 + i * 1.88, 5.95, c)
    footer(s)


def slide_7(prs, assets):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s); title(s, 7, "目标人群：按参与动机设计内容，而不是让所有人排同一条队", "不同人群对健康的关注点不同，现场应通过签到分流和内容推荐提高停留质量。", "人群")
    personas = [
        ("周边社区居民", "动机：家门口、低门槛、可信赖。\n匹配：基础体检、健康手册、医生咨询。", COL["blue"]),
        ("中老年群体", "动机：血压血糖、慢病风险、运动安全。\n匹配：慢病咨询、八段锦、用药护理提示。", COL["teal"]),
        ("职场人群", "动机：亚健康、体态、压力、睡眠。\n匹配：BMI、心率、肩颈拉伸、体态评估。", COL["green"]),
        ("亲子家庭", "动机：家庭健康习惯和儿童健康认知。\n匹配：健康问答、饮食贴纸、亲子打卡。", COL["orange"]),
        ("运动关注者", "动机：科学运动和身体管理。\n匹配：拉伸互动、健步走指导、运动康复建议。", COL["red"]),
    ]
    for i, (h, b, c) in enumerate(personas):
        card(s, h, b, 0.72 + (i % 3) * 4.05, 1.28 + (i // 3) * 1.56, 3.65, 1.2, c, 8.4)
    picture(s, assets["exercise"], 8.88, 4.34, 3.32, 1.82)
    card(s, "现场分流建议", "签到区用不同颜色的参与卡快速识别需求：体检卡、慢病咨询卡、职场拉伸卡、亲子问答卡、运动打卡卡。志愿者按卡片推荐路线，避免体检区瞬时拥堵。", 0.72, 4.82, 7.65, 1.32, COL["teal"], 9.2, COL["mint"])
    footer(s)


def slide_8(prs, assets):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s); title(s, 8, "主题创意：被守护的健康，守护者的节日", "主推主题兼具护士节情感、公益属性和市民参与感，适合医院公益健康活动传播。", "主题")
    picture(s, assets["tribute"], 0.75, 1.28, 4.65, 3.0)
    textbox(s, "主推活动主题", 5.78, 1.34, 2.0, 0.24, 10, COL["teal"], True)
    textbox(s, "致敬白衣天使\n共赴健康生活", 5.78, 1.72, 5.95, 0.98, 24, COL["navy"], True, spacing=0.9)
    pill(s, "传播口号：今天把掌声献给守护者，也把健康带回家", 5.8, 2.95, 5.55, 0.42, COL["blue"], size=9)
    textbox(
        s,
        "主题释义：护士是健康守护链条中最温暖、最贴近患者的一环。护士节当天，医院把这份守护延伸到市民身边，"
        "用公益体检、专业咨询和健康生活互动，让“被守护的健康”变成可以参与、可以感受、可以分享的行动。",
        5.8, 3.55, 5.9, 0.95, 11.2, COL["gray"],
    )
    backups = [
        ("守护在身边，健康向未来", "一场公益体检，开启一段健康陪伴。"),
        ("天使守护日，市民健康行", "走进公益健康现场，遇见更懂自己的健康。"),
        ("以爱致敬，以健康同行", "向护理团队致敬，与科学生活同行。"),
    ]
    for i, (h, b) in enumerate(backups):
        card(s, h, b, 0.75 + i * 4.05, 5.18, 3.65, 0.92, [COL["blue"], COL["green"], COL["orange"]][i], 8)
    footer(s)


def slide_9(prs, assets):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s); title(s, 9, "活动亮点：让医院专业服务成为市民可感知的现场体验", "每个亮点都对应一个现场动作，既服务公益价值，也服务传播画面和后续连接。", "亮点")
    highlights = [
        ("护士节致敬仪式", "通过短片、鲜花、证书与公众祝福墙，正式表达对护理团队的尊重。", COL["blue"]),
        ("市民公益体检", "开放血压、血糖、BMI、心率等基础检测，让市民即时参与。", COL["teal"]),
        ("专业健康咨询", "医生、护士面对面解释指标，提升医院专业可信度。", COL["green"]),
        ("运动健康互动", "用拉伸、八段锦、健步走和体态评估让科学运动变成现场动作。", COL["orange"]),
        ("健康生活打卡", "设置健康承诺卡和轻量挑战，鼓励市民把健康行动带回家。", COL["red"]),
        ("朋友圈传播", "打卡墙、手举牌和亲和文案，让真实参与自然进入社交平台。", COL["blue2"]),
        ("医院服务展示", "以公益方式展示体检、慢病管理、护理门诊和社区服务能力。", COL["teal"]),
    ]
    for i, (h, b, c) in enumerate(highlights):
        x = 0.7 + (i % 4) * 3.05
        y = 1.28 + (i // 4) * 1.64
        card(s, h, b, x, y, 2.75, 1.25, c, 7.7)
    picture(s, assets["checkup"], 0.72, 4.88, 3.35, 1.7)
    picture(s, assets["exercise"], 4.3, 4.88, 3.35, 1.7)
    picture(s, assets["phone"], 8.05, 4.88, 2.03, 1.7)
    card(s, "亮点落地提醒", "所有亮点都需要明确负责人、物料和完成标准：仪式有流程单，体检有排队规则，互动有安全提示，传播有指定拍摄点。", 10.35, 4.88, 2.08, 1.7, COL["blue"], 7.6, COL["mint"])
    footer(s)


def slide_10(prs, assets):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s); title(s, 10, "整体活动架构：七大区域承接完整参与路径", "市民从签到进入，到体检、咨询、运动、科普、打卡和后续服务，形成连续体验。", "架构")
    picture(s, assets["map"], 0.72, 1.2, 5.75, 3.72)
    zones = [
        ("仪式区", "开场、领导致辞、护士代表分享、短片播放和集体合影。"),
        ("公益体检区", "信息登记、基础检测、结果记录和异常指标引导。"),
        ("健康咨询区", "内科、护理、慢病、运动康复和营养生活方式咨询。"),
        ("运动互动区", "拉伸教学、八段锦、健步走指导、体态评估和打卡挑战。"),
        ("科普展示区", "体检、运动、慢病、膳食、睡眠五类展板与资料领取。"),
        ("拍照打卡区", "护士节致敬墙、祝福墙、健康承诺卡和朋友圈手举牌。"),
        ("报名转化区", "后续体检、慢病管理、护理门诊和社区服务二维码咨询。"),
    ]
    for i, (h, b) in enumerate(zones):
        x = 6.75 + (i % 2) * 2.9
        y = 1.24 + (i // 2) * 1.04
        card(s, h, b, x, y, 2.6, 0.78, [COL["blue"], COL["teal"], COL["green"], COL["orange"]][i % 4], 6.9)
    card(s, "动线原则", "入口不堵门，体检区不压主通道，咨询区靠近体检出口但保留私密感，打卡区靠近离场路径，服务二维码放在市民体验完成之后。", 0.72, 5.24, 11.7, 0.9, COL["teal"], 9.4, COL["mint"])
    footer(s)


def slide_11(prs, assets):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s); title(s, 11, "半天流程规划：仪式短、体验足、咨询细、传播自然", "建议按 08:30-12:00 组织，保证仪式正式感，同时留出充足时间给公益体检与咨询。", "流程")
    steps = [
        ("08:30", "签到入场", "领取号码贴、体检流程卡与健康手册，志愿者按需求推荐路线。", COL["blue"]),
        ("09:00", "主持开场", "介绍护士节意义、公益属性、现场规则和老人优先安排。", COL["teal"]),
        ("09:10", "护士致敬", "领导致辞、护士代表分享、致敬短片、献花证书和合影。", COL["green"]),
        ("09:30", "体检咨询", "分批进行血压、血糖、BMI、心率、问诊和指标解读。", COL["orange"]),
        ("10:40", "健康讲座", "用 20 分钟讲清定期体检、慢病预防或科学运动主题。", COL["red"]),
        ("11:10", "运动互动", "开展拉伸教学、八段锦、体态评估和健康打卡挑战。", COL["blue2"]),
        ("11:40", "合影离场", "祝福墙合影、朋友圈打卡、扫码领取资料并引导后续服务。", COL["teal"]),
    ]
    y = 1.35
    for i, (time, h, b, c) in enumerate(steps):
        x = 0.78 + i * 1.78
        icon_circle(s, str(i + 1), x + 0.32, y, c)
        textbox(s, time, x, y + 0.62, 1.08, 0.22, 10.5, c, True, PP_ALIGN.CENTER)
        textbox(s, h, x - 0.07, y + 0.92, 1.22, 0.24, 9.5, COL["navy"], True, PP_ALIGN.CENTER)
        textbox(s, b, x - 0.18, y + 1.22, 1.55, 0.95, 7.2, COL["gray"], False, PP_ALIGN.CENTER)
        if i < len(steps) - 1:
            arrow = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x + 1.12), Inches(y + 0.13), Inches(0.48), Inches(0.18))
            fill(arrow, COL["light_gray"]); no_line(arrow)
    picture(s, assets["tribute"], 0.78, 4.2, 3.65, 1.75)
    picture(s, assets["checkup"], 4.75, 4.2, 3.65, 1.75)
    card(s, "流程控制点", "09:30 后要快速分流，避免所有市民同时进入体检区。主持人需反复提醒排队顺序、空腹/餐后血糖差异、老人优先和检测隐私保护。", 8.74, 4.2, 3.68, 1.75, COL["blue"], 8.6, COL["mint"])
    footer(s)


def slide_12(prs, assets):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s); title(s, 12, "现场区域规划：把排队、隐私、效率、传播和安全一起设计", "平面布局不是简单摆桌子，而是服务市民动线、医护效率和现场秩序。", "现场")
    picture(s, assets["map"], 0.65, 1.18, 5.05, 3.36)
    rules = [
        ("排队效率", "入口设置签到和号码贴，热门体检项目配置等候线、座椅和分流提示。"),
        ("体检隐私", "检测结果不公开播报，体检台与咨询台保持距离，必要时配置屏风。"),
        ("咨询效率", "咨询区紧邻体检出口，志愿者提前帮市民整理检测结果和问题。"),
        ("拍照传播", "打卡区靠近出口，完成体验后自然进入合影和朋友圈分享环节。"),
        ("安全秩序", "消防通道不摆物料，线缆入线槽，运动区设置防滑垫和安全距离。"),
    ]
    for i, (h, b) in enumerate(rules):
        card(s, h, b, 6.05 + (i % 2) * 3.1, 1.18 + (i // 2) * 1.08, 2.8, 0.82, [COL["blue"], COL["teal"], COL["green"], COL["orange"], COL["red"]][i], 7.3)
    card(s, "进场前确认清单", "承办方需提前测量场地尺寸、电源点位、消防通道、主入口、媒体机位和雨天备选位置；最终点位图需经医院场地、安保和医护负责人共同确认。", 0.65, 5.0, 11.75, 0.95, COL["teal"], 9.4, COL["mint"])
    footer(s)


def slide_13(prs, assets):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s); title(s, 13, "公益体检区：让市民测得到、看得懂、问得清", "体检区不是简单测量数据，而是完成“检测结果 → 初步解释 → 咨询引导”的公益服务闭环。", "体检")
    picture(s, assets["checkup"], 0.68, 1.22, 4.5, 3.22)
    card(s, "建议体检项目", "血压测量、指尖血糖、身高体重与 BMI、静息心率、基础问诊、慢病风险初筛和护理健康建议。项目以医院医护最终确认的服务能力为准。", 5.55, 1.22, 3.2, 1.28, COL["blue"], 8.7)
    card(s, "市民体验流程", "登记领取号码贴 → 按号码进入检测 → 记录结果 → 护士初步说明 → 异常指标引导至医生咨询台或后续检查入口。", 9.12, 1.22, 3.2, 1.28, COL["teal"], 8.7)
    card(s, "桌台配置", "每个项目设置独立桌牌、检测设备、一次性耗材、登记表、酒精棉片、手消用品、废弃物收纳和隐私提示牌。", 5.55, 2.88, 3.2, 1.28, COL["green"], 8.7)
    card(s, "秩序与隐私", "等候区设置座椅，老人和行动不便者优先。检测结果不公开播报，记录表由市民本人保管，异常情况由医护人员低声沟通。", 9.12, 2.88, 3.2, 1.28, COL["orange"], 8.7)
    card(s, "执行边界", "医疗检测和健康判断由医院医护人员主导；承办方负责动线、桌椅、排队、耗材补给、秩序维护、导视物料和信息流转，不替代医疗判断。", 0.68, 5.05, 11.65, 0.9, COL["teal"], 9.2, COL["mint"])
    footer(s)


def slide_14(prs, assets):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s); title(s, 14, "健康咨询区：把专业服务讲得亲切、准确、可执行", "咨询区是医院专业度和公益服务温度的集中体现，要让市民带着问题来，带着建议走。", "咨询")
    picture(s, assets["consult"], 7.62, 1.18, 4.62, 3.1)
    desks = [
        ("内科基础咨询", "解释血压、血糖、心率等基础指标，提示复查和就医路径。"),
        ("护理健康指导", "围绕居家护理、用药习惯、基础照护和健康观察给出建议。"),
        ("慢病管理咨询", "面向高血压、糖尿病等风险人群，提供长期监测和管理提示。"),
        ("运动康复建议", "针对久坐、肩颈腰背不适和运动安全，提供轻量动作建议。"),
        ("营养生活方式", "围绕控盐控糖、饮食结构、睡眠压力和体重管理提出可执行行动。"),
    ]
    for i, (h, b) in enumerate(desks):
        card(s, h, b, 0.72 + (i % 2) * 3.38, 1.2 + (i // 2) * 1.08, 3.05, 0.82, [COL["blue"], COL["teal"], COL["green"], COL["orange"], COL["red"]][i], 7.2)
    card(s, "沟通话术原则", "现场建议使用“建议您关注”“可以先从这三件事开始”“如持续异常建议到院进一步检查”等表达，既保持专业边界，也避免给市民造成不必要焦虑。", 0.72, 4.75, 6.55, 1.02, COL["teal"], 9.2, COL["mint"])
    card(s, "转化方式", "咨询结束后由志愿者引导扫码了解医院体检中心、慢病管理门诊、护理门诊或社区健康服务，保持公益语境，不做强推销售。", 7.62, 4.75, 4.62, 1.02, COL["blue"], 9.2, COL["white"])
    footer(s)


def slide_15(prs, assets):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s); title(s, 15, "运动健康互动区：让科学运动变成市民会做的动作", "互动内容必须轻量、安全、容易跟练，适合普通市民而不是专业运动人群。", "互动")
    picture(s, assets["exercise"], 0.72, 1.18, 5.2, 3.35)
    activities = [
        ("科学拉伸教学", "肩颈、腰背、下肢拉伸，每轮 5 分钟，适合职场人和中老年。"),
        ("八段锦/健步走", "动作慢、示范清楚，强调循序渐进和身体感受，不追求强度。"),
        ("体态评估体验", "通过站姿观察和肩颈活动度测试，引导久坐人群关注身体信号。"),
        ("运动打卡挑战", "完成一个拉伸动作或健康步数承诺，即可获得贴纸或健康小礼。"),
    ]
    for i, (h, b) in enumerate(activities):
        card(s, h, b, 6.25 + (i % 2) * 3.05, 1.2 + (i // 2) * 1.26, 2.75, 0.96, [COL["blue"], COL["teal"], COL["green"], COL["orange"]][i], 7.4)
    card(s, "现场安全要求", "运动区地面使用防滑垫，周边保留安全距离；不安排高强度跳跃，不要求老人勉强完成动作。现场配备休息座椅、饮用水和医护应急支持。", 0.72, 5.03, 6.1, 0.92, COL["teal"], 9.2, COL["mint"])
    card(s, "互动节奏", "每轮互动控制在 8-10 分钟，主持人负责召集，指导员负责示范，志愿者负责安全提醒、奖品发放和打卡拍照。", 7.12, 5.03, 5.1, 0.92, COL["blue"], 9.2, COL["white"])
    footer(s)


def slide_16(prs, assets):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s); title(s, 16, "护士节致敬环节：正式、温暖、不煽情过度", "致敬环节要让护理团队被看见、被尊重，也让市民理解护理工作背后的专业价值。", "仪式")
    picture(s, assets["tribute"], 0.72, 1.18, 5.05, 3.4)
    steps = [
        ("主持开场", "用“今天，我们把掌声献给每一位守护生命细节的人”引入活动。"),
        ("护士代表分享", "邀请 1-2 位护士讲述真实护理片段，每人控制在 3 分钟内。"),
        ("致敬短片", "展示护理日常、夜班值守、患者沟通和公益服务画面。"),
        ("鲜花/证书", "由医院领导或市民代表向护士代表献花，颁发纪念证书。"),
        ("公众祝福墙", "市民写下想对护士说的话，并与护士团队合影。"),
    ]
    for i, (h, b) in enumerate(steps):
        card(s, h, b, 6.08 + (i % 2) * 3.12, 1.2 + (i // 2) * 1.02, 2.82, 0.76, [COL["blue"], COL["teal"], COL["green"], COL["orange"], COL["red"]][i], 7.1)
    card(s, "彩排与传播重点", "仪式建议控制在 20 分钟以内。短片、音乐、证书名单、上台顺序和合影站位需提前彩排；摄影摄像必须捕捉献花、合影、祝福墙和护理团队特写。", 0.72, 5.08, 11.72, 0.88, COL["teal"], 9.1, COL["mint"])
    footer(s)


def slide_17(prs, assets):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s); title(s, 17, "健康生活科普：把专业知识变成一句能记住的行动建议", "科普内容围绕体检、运动、慢病、膳食、睡眠五个高频话题展开，适合展板、讲座和手册复用。", "科普")
    picture(s, assets["education"], 0.72, 1.14, 4.7, 3.1)
    topics = [
        ("定期体检", "体检不是“有病才做”，而是帮助我们更早发现身体发出的提醒。今日行动：记录一次血压、血糖或 BMI。"),
        ("科学运动", "适合自己的运动，应该能长期坚持、循序渐进，并在运动后感觉舒展而不是透支。今日行动：学习一个拉伸动作。"),
        ("慢病预防", "血压、血糖、血脂的变化往往从生活习惯开始，也需要通过持续监测来管理。今日行动：建立家庭健康记录。"),
        ("合理膳食", "少油、少盐、少糖可以从每餐多一份蔬菜、少一杯含糖饮料开始。今日行动：给晚餐减一勺盐。"),
        ("睡眠压力", "睡不好、压力大也会影响血压、体重和免疫状态。今日行动：今晚提前 30 分钟放下手机。"),
    ]
    for i, (h, b) in enumerate(topics):
        card(s, h, b, 5.78 + (i % 2) * 3.15, 1.12 + (i // 2) * 1.15, 2.85, 0.92, [COL["blue"], COL["teal"], COL["green"], COL["orange"], COL["red"]][i], 6.8)
    card(s, "内容审核机制", "所有科普文案先由承办方转化为通俗表达，再提交医院相关科室审核，确保专业准确、表述温和、适合公众阅读。", 0.72, 5.03, 11.7, 0.9, COL["teal"], 9.2, COL["mint"])
    footer(s)


def slide_18(prs, assets):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s); title(s, 18, "朋友圈宣发与引流：让真实参与成为自然传播素材", "文案要亲和、有行动指令、适合市民转发，避免过度官方和营销化。", "传播")
    picture(s, assets["phone"], 0.78, 1.2, 2.9, 4.55)
    copies = [
        ("预热文案", "这个护士节，医院把公益体检和健康咨询带到大家身边。血压、血糖、BMI、运动指导都能现场参与，带上家人一起来给健康做个小检查吧。"),
        ("活动当天", "今天在医院公益健康活动现场，测了血压，也学了几个肩颈拉伸动作。谢谢护士们的耐心讲解，健康真的要从日常开始。"),
        ("现场打卡", "我在“致敬白衣天使，共赴健康生活”现场完成健康打卡。把掌声送给守护者，也把健康提醒送给自己和家人。"),
        ("活动复盘", "一上午的公益体检和健康咨询结束，感谢每一位到场市民，也感谢护理团队的专业守护。愿我们都从今天开始更关注身体发出的信号。"),
    ]
    for i, (h, b) in enumerate(copies):
        card(s, h, b, 4.12 + (i % 2) * 4.05, 1.18 + (i // 2) * 1.78, 3.65, 1.42, [COL["blue"], COL["teal"], COL["green"], COL["orange"]][i], 8.1)
    card(s, "拍摄与发布建议", "摄影摄像提前拿到必拍清单：护士致敬、体检服务、咨询沟通、运动互动、打卡合影、物料全景。活动中快速筛图，供医院公众号和朋友圈当天发布。", 4.12, 5.24, 7.7, 0.86, COL["teal"], 9.1, COL["mint"])
    footer(s)


def slide_19(prs, assets):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s); title(s, 19, "视觉与美陈：干净专业，温暖亲民，适合现场出片", "视觉系统以医疗蓝、白色、浅绿色为主，辅以少量暖色表达节日温度。", "视觉")
    # palette
    textbox(s, "主视觉延展", 0.72, 1.14, 1.4, 0.24, 10, COL["teal"], True)
    for i, (name, c) in enumerate([("医疗蓝", COL["blue"]), ("洁净白", COL["white"]), ("浅绿", COL["teal"]), ("暖橙", COL["orange"]), ("柔粉", COL["pink"])]):
        rect = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.75 + i * 0.9), Inches(1.55), Inches(0.65), Inches(0.65))
        fill(rect, c); line_color(rect, COL["light_gray"], 0.6)
        textbox(s, name, 0.65 + i * 0.9, 2.28, 0.85, 0.22, 6.6, COL["gray"], False, PP_ALIGN.CENTER)
    picture(s, assets["stage"], 6.75, 1.12, 5.35, 2.8)
    elements = [
        ("签到背景板", "标题清晰、LOGO规范、保留合影空间，适合媒体和市民拍照。"),
        ("护士节致敬墙", "设置祝福便利贴和护士团队照片，让市民表达感谢。"),
        ("健康打卡装置", "使用“今日健康承诺”“把健康带回家”等文案形成朋友圈记忆点。"),
        ("体检区导视", "用颜色区分血压、血糖、BMI、咨询出口，减少现场口头解释。"),
        ("咨询区桌牌", "统一科室名称、咨询范围、医护姓名和二维码，强化专业秩序。"),
    ]
    for i, (h, b) in enumerate(elements):
        card(s, h, b, 0.72 + (i % 2) * 3.05, 3.15 + (i // 2) * 1.0, 2.76, 0.75, [COL["blue"], COL["teal"], COL["green"], COL["orange"], COL["red"]][i], 6.8)
    card(s, "一体化落地能力", "承办方负责视觉设计、喷绘制作、结构搭建、现场安装和撤场回收。所有含医院 LOGO 和专业科普内容的物料需提前提交甲方审核。", 6.75, 4.32, 5.35, 1.02, COL["blue"], 9.2, COL["mint"])
    footer(s)


def slide_20(prs, assets):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s); title(s, 20, "舞美与物料搭建：自有团队保障进撤场效率与现场稳定", "依托自有舞美搭建执行团队和自有舞台活动设备，减少外协沟通成本，提高现场响应速度。", "搭建")
    picture(s, assets["stage"], 0.72, 1.18, 5.35, 3.28)
    groups = [
        ("舞台区", "小型主舞台或平面仪式区、主背景板、讲台、主持位、合影站位、摄影通道。"),
        ("显示系统", "LED 屏或电视屏，用于播放致敬短片、流程提示、健康科普和现场照片。"),
        ("灯光音响", "主扩音响、无线麦、补光灯和基础氛围灯，保证主持、分享和讲座清晰可听。"),
        ("签到导视", "签到台、欢迎立牌、总导览牌、区域导视、排队地贴和咨询引导牌。"),
        ("展陈物料", "科普展板、易拉宝、桌牌、手举牌、祝福墙、打卡装置、健康承诺卡。"),
        ("现场保障", "线槽、配电箱、工具包、备用麦、电池、雨棚、防滑垫和应急物料。"),
    ]
    for i, (h, b) in enumerate(groups):
        card(s, h, b, 6.38 + (i % 2) * 3.05, 1.16 + (i // 2) * 1.18, 2.75, 0.9, [COL["blue"], COL["teal"], COL["green"], COL["orange"], COL["red"], COL["blue2"]][i], 7.1)
    card(s, "搭建节奏", "建议活动前一天或前半天进场，完成结构搭建、画面安装、电源测试、灯光音响联调和安全检查；活动结束后按医院要求撤场，保证场地恢复整洁。", 0.72, 5.25, 11.62, 0.85, COL["teal"], 9.2, COL["mint"])
    footer(s)


def slide_21(prs, assets):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s); title(s, 21, "人员分工与执行保障：职责清晰，现场才稳定", "医院医护、承办执行、志愿者和安保需要提前定责、定岗、定流程。", "执行")
    picture(s, assets["team"], 0.72, 1.2, 4.5, 3.12)
    roles = [
        ("项目总控", "统筹进度、现场决策、甲方沟通和突发问题处理。"),
        ("客户对接", "确认医院需求、物料审核、医护排班和流程节点。"),
        ("设计团队", "完成主视觉、物料延展、导视系统和朋友圈海报。"),
        ("搭建设备", "负责舞台、美陈、灯光音响、显示设备、桌椅和电源安全。"),
        ("主持摄影", "把控仪式节奏、讲座衔接、活动记录和传播素材采集。"),
        ("医护协调", "由医院安排检测和咨询人员，明确项目边界、耗材和隐私要求。"),
        ("志愿秩序", "负责签到、分流、排队、老人协助、物料发放和打卡引导。"),
        ("物料管理", "建立物料清单、进场验收、现场补给和撤场回收机制。"),
    ]
    for i, (h, b) in enumerate(roles):
        card(s, h, b, 5.55 + (i % 2) * 3.35, 1.12 + (i // 2) * 1.0, 3.05, 0.76, [COL["blue"], COL["teal"], COL["green"], COL["orange"]][i % 4], 6.6)
    card(s, "执行交底", "活动前至少进行一次全员线上确认和一次现场走台，重点确认上台顺序、检测流程、应急联系人、设备开关机、摄影点位和撤场路径。", 0.72, 5.0, 11.66, 0.9, COL["teal"], 9.2, COL["mint"])
    footer(s)


def slide_22(prs, assets):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s); title(s, 22, "风险预案：把公益现场的不确定性前置管理", "风险管理覆盖人流、排队、天气、设备、老人安全、医疗隐私和突发不适。", "风险")
    s.shapes.add_picture(str(assets["team"]), Inches(11.72), Inches(0.72), Inches(0.68), Inches(0.43))
    risks = [
        ("人流拥堵", "分时段预约与现场号码贴并行；入口设置分流志愿者，热门项目增加等候座椅。"),
        ("体检排队", "每个项目设置等待提示，体检区满负荷时引导市民先参与科普或运动互动。"),
        ("天气变化", "户外准备帐篷、雨棚、防滑垫和设备防雨方案；高温增加饮水点和休息区。"),
        ("设备故障", "音响、麦克风、显示屏、血压计等准备备用设备和电池，技术人员全程值守。"),
        ("老人安全", "运动前提醒量力而行，出现头晕、胸闷等不适立即停止并联系现场医护。"),
        ("医疗隐私", "检测结果不公开展示，登记表不随意堆放，咨询区保持适当沟通距离。"),
        ("现场秩序", "安保与志愿者保持通道畅通，消防通道不摆物料，线缆使用线槽固定。"),
        ("突发不适", "由医院医护按医疗流程处置，承办方负责疏导围观、保留通道和通知总控。"),
    ]
    for i, (h, b) in enumerate(risks):
        card(s, h, b, 0.72 + (i % 2) * 6.05, 1.12 + (i // 2) * 1.12, 5.58, 0.86, [COL["blue"], COL["teal"], COL["green"], COL["orange"], COL["red"], COL["blue2"], COL["teal"], COL["orange"]][i], 7.4)
    card(s, "响应机制", "现场统一采用“谁发现、谁上报、谁处理、谁对外沟通”的链路。医疗相关事件以医院医护判断和流程为准，承办方配合秩序维护与现场沟通。", 0.72, 5.93, 11.65, 0.72, COL["teal"], 8.8, COL["mint"])
    footer(s)


def slide_23(prs, assets):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s); title(s, 23, "预期效果：让一次公益活动沉淀为医院长期连接", "活动成果从品牌、参与、公益、传播、社区关系和后续健康管理六个维度评估。", "成果")
    picture(s, assets["metrics"], 0.75, 1.16, 4.75, 3.1)
    outcomes = [
        ("医院品牌", "强化医院专业、可靠、有温度的公众形象，让护理团队的守护精神被看见。"),
        ("公众参与", "让市民至少完成一次检测、一次咨询或一次运动体验，提升健康行动意愿。"),
        ("公益影响", "把医院服务主动送到市民身边，形成社区和媒体可记录的公益案例。"),
        ("传播曝光", "通过现场照片、朋友圈打卡、医院公众号和社区群形成二次传播。"),
        ("社区关系", "增强医院与周边居民、社区组织、家庭人群之间的信任和互动。"),
        ("后续转化", "为体检中心、慢病管理、护理门诊和健康管理服务建立可信认知入口。"),
    ]
    for i, (h, b) in enumerate(outcomes):
        card(s, h, b, 5.85 + (i % 2) * 3.15, 1.14 + (i // 2) * 1.16, 2.85, 0.9, [COL["blue"], COL["teal"], COL["green"], COL["orange"], COL["red"], COL["blue2"]][i], 7.2)
    card(s, "复盘交付", "活动结束后 24 小时内提交快报：精选照片、核心数据和传播建议；3 个工作日内提交完整复盘，包含问题、改进和后续服务建议。", 0.75, 5.08, 11.55, 0.86, COL["teal"], 9.2, COL["mint"])
    footer(s)


def slide_24(prs, assets):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    s.shapes.add_picture(str(assets["cover"]), Inches(6.65), Inches(0.82), Inches(5.72), Inches(3.45))
    textbox(s, "以专业守护健康\n以温度连接市民", 0.82, 1.34, 5.9, 1.05, 27, COL["navy"], True, spacing=0.9)
    textbox(
        s,
        "让护士节不止是一场致敬，更成为医院走近社区、服务公众、倡导健康生活的一次温暖行动。",
        0.86, 2.78, 5.62, 0.58, 14, COL["blue"], True,
    )
    textbox(
        s,
        "我们将以一体化活动策划与执行能力，完成从主题创意、视觉设计、美陈舞台搭建、设备保障、人员统筹到现场复盘的全流程落地，"
        "协助医院把公益服务做得有秩序、有温度、有传播力。",
        0.86, 3.58, 5.7, 1.0, 11.5, COL["gray"],
    )
    card(s, "汇报收束", "本方案的核心价值，是把护理团队的节日情感、医院公益服务和市民健康行动连接在同一个现场。", 0.86, 5.04, 5.7, 0.92, COL["teal"], 9.8, COL["mint"])
    pill(s, "THANK YOU", 0.86, 6.28, 1.65, 0.36, COL["blue"], size=9.2)
    textbox(s, "医院护士节公益健康活动方案｜直接汇报版", 0.86, 6.82, 4.4, 0.2, 8.2, COL["gray"])


SLIDE_FUNCS = [
    slide_1, slide_2, slide_3, slide_4, slide_5, slide_6, slide_7, slide_8,
    slide_9, slide_10, slide_11, slide_12, slide_13, slide_14, slide_15, slide_16,
    slide_17, slide_18, slide_19, slide_20, slide_21, slide_22, slide_23, slide_24,
]


def build():
    assets = create_assets()
    prs = Presentation()
    prs.slide_width = Inches(WIDE_W)
    prs.slide_height = Inches(WIDE_H)
    for fn in SLIDE_FUNCS:
        fn(prs, assets)
    prs.core_properties.title = "护士节公益健康活动方案_直接汇报版"
    prs.core_properties.subject = "医院护士节公益体检与健康生活倡导活动提案"
    prs.core_properties.author = "Codex"
    prs.save(OUT)
    print(f"saved report pptx with {len(prs.slides)} slides")


if __name__ == "__main__":
    build()
